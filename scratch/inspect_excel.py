import openpyxl
import os

file_path = r"d:\расм\public\1-brigada.xlsx"

if not os.path.exists(file_path):
    print("File not found:", file_path)
    exit(1)

wb = openpyxl.load_workbook(file_path, data_only=False)
sheet1 = wb.worksheets[1]

print("--- worker row 11 and 12 cell formulas ---")
for col in range(1, 23):
    cell = sheet1.cell(row=11, column=col)
    cell_ref = f"{openpyxl.utils.get_column_letter(col)}11"
    print(f"  {cell_ref}: Value/Formula = {cell.value}")

for col in range(1, 23):
    cell = sheet1.cell(row=12, column=col)
    cell_ref = f"{openpyxl.utils.get_column_letter(col)}12"
    print(f"  {cell_ref}: Value/Formula = {cell.value}")

wb.close()
