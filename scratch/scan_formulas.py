import openpyxl
import os

file_path = r"d:\расм\public\1-brigada.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=False)
sheet = wb.worksheets[1]

formulas = {}
for r in range(1, sheet.max_row + 1):
    for c in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=r, column=c)
        if cell.value and str(cell.value).startswith('='):
            cell_ref = f"{openpyxl.utils.get_column_letter(c)}{r}"
            formulas[cell_ref] = cell.value

print(f"Total formulas found in sheet 1: {len(formulas)}")
print("\nFirst 30 formulas:")
for k, v in list(formulas.items())[:30]:
    print(f"  {k}: {v}")

print("\nFormulas containing spaces:")
for k, v in formulas.items():
    if ' ' in str(v):
        print(f"  {k}: {v}")

wb.close()
