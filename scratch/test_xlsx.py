import openpyxl
import os

file_path = r"d:\расм\public\1-brigada.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
sheet = wb.worksheets[1]

print("--- Row 11 cell values & types ---")
for col in range(3, 19): # C to R
    cell = sheet.cell(row=11, column=col)
    cell_ref = f"{openpyxl.utils.get_column_letter(col)}11"
    print(f"  {cell_ref}: {cell.value} (Type: {type(cell.value).__name__})")

print("\n--- Row 12 cell values & types ---")
for col in range(3, 19): # C to R
    cell = sheet.cell(row=12, column=col)
    cell_ref = f"{openpyxl.utils.get_column_letter(col)}12"
    print(f"  {cell_ref}: {cell.value} (Type: {type(cell.value).__name__})")

print("\n--- Cell S11 (SUM) ---")
cell_s11 = sheet.cell(row=11, column=19)
print(f"  S11 value: {cell_s11.value} (Type: {type(cell_s11.value).__name__})")

wb.close()
