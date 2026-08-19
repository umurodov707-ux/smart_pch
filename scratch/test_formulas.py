import openpyxl
import os
import re

file_path = r"d:\расм\public\1-brigada.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=False)
sheet = wb.worksheets[1]

# Replicate pu74FormulaMap
pu74FormulaMap = {}
for r in range(1, sheet.max_row + 1):
    for c in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=r, column=c)
        if cell.value and str(cell.value).startswith('='):
            cell_ref = f"{openpyxl.utils.get_column_letter(c)}{r}"
            pu74FormulaMap[cell_ref] = cell.value

# Replicate pu74EditedCells
pu74EditedCells = {}

def getCellNumericValue(cell_ref):
    if cell_ref in pu74EditedCells:
        try:
            return float(pu74EditedCells[cell_ref])
        except ValueError:
            return 0.0

    if cell_ref in pu74FormulaMap:
        result = evaluateFormula(pu74FormulaMap[cell_ref])
        if result is not None:
            return result

    # Read from cell
    cell = sheet[cell_ref]
    # We open workbook with data_only=False, so to get value we might need to look at a data_only workbook
    # But since we simulate formulas, we can load a data_only workbook for fallbacks
    return get_fallback_value(cell_ref)

# Load data_only workbook for cell value fallbacks
wb_data = openpyxl.load_workbook(file_path, data_only=True)
sheet_data = wb_data.worksheets[1]

def get_fallback_value(cell_ref):
    val = sheet_data[cell_ref].value
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0

def evaluateFormula(formula):
    expr = str(formula).strip()
    if expr.startswith('='):
        expr = expr[1:]
    if expr.startswith('+'):
        expr = expr[1:]

    # 1. SUM range
    def sum_range_sub(match):
        start_ref = match.group(1)
        end_ref = match.group(2)
        
        # Decode cell refs
        start_c = openpyxl.utils.column_index_from_string(re.match(r"[A-Z]+", start_ref).group())
        start_r = int(re.search(r"\d+", start_ref).group())
        end_c = openpyxl.utils.column_index_from_string(re.match(r"[A-Z]+", end_ref).group())
        end_r = int(re.search(r"\d+", end_ref).group())
        
        total = 0.0
        min_row = min(start_r, end_r)
        max_row = max(start_r, end_r)
        min_col = min(start_c, end_c)
        max_col = max(start_c, end_c)
        
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                ref = f"{openpyxl.utils.get_column_letter(c)}{r}"
                total += getCellNumericValue(ref)
        return str(total)

    expr = re.sub(r"SUM\((?:'[A-Za-z0-9_-]+'!)?([\$]?[A-Z]+[\$]?\d+):([\$]?[A-Z]+[\$]?\d+)\)", sum_range_sub, expr, flags=re.IGNORECASE)

    # 2. SUM list
    def sum_list_sub(match):
        list_str = match.group(1)
        refs = list_str.split(',')
        total = 0.0
        for ref in refs:
            clean_ref = ref.strip()
            # Check if reference
            if re.match(r"^\$?[A-Z]+\$?\d+$", clean_ref, re.IGNORECASE):
                total += getCellNumericValue(clean_ref)
            else:
                try:
                    total += float(clean_ref)
                except ValueError:
                    pass
        return str(total)

    expr = re.sub(r"SUM\(([^)]+)\)", sum_list_sub, expr, flags=re.IGNORECASE)

    # 3. Replace individual cells
    def cell_sub(match):
        cell_ref = match.group(0)
        return str(getCellNumericValue(cell_ref))

    expr = re.sub(r"\$?[A-Z]+\$?\d+", cell_sub, expr, flags=re.IGNORECASE)

    # 4. Safe eval
    # Clean expression of anything other than numbers, math operators
    expr_clean = re.sub(r"[^0-9\+\-\*\/\(\)\.]", "", expr)
    if expr_clean:
        try:
            return eval(expr_clean)
        except Exception as e:
            print(f"Error evaluating clean expr: {expr_clean} from original {formula} - {e}")
            return 0.0
    return 0.0

print("--- Testing individual worker S cells ---")
for r in range(11, 21):
    ref = f"S{r}"
    if ref in pu74FormulaMap:
        print(f"Cell {ref} formula: {pu74FormulaMap[ref]} -> Evaluated: {getCellNumericValue(ref)}")

print("\n--- Testing Q41, R41, S41 ---")
print("Q41 formula:", pu74FormulaMap.get('Q41'), "-> Evaluated:", getCellNumericValue('Q41'))
print("R41 formula:", pu74FormulaMap.get('R41'), "-> Evaluated:", getCellNumericValue('R41'))
print("S41 formula:", pu74FormulaMap.get('S41'), "-> Evaluated:", getCellNumericValue('S41'))

print("\n--- Testing S43 (SUM S11:S40) ---")
print("S43 formula:", pu74FormulaMap.get('S43'), "-> Evaluated:", getCellNumericValue('S43'))

wb.close()
wb_data.close()
